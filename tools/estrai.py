#!/usr/bin/env python3
# estrai.py — QuizNCC
#
# Legge l'xlsm della Camera di Commercio di Milano (quiz iscrizione al ruolo
# conducenti taxi/NCC, ed. 22/11/2019) e produce i JSON che la webapp carica.
# Si lancia una volta; se l'xlsm cambia, si rilancia.
#
#   python3 tools/estrai.py tools/quiz-ccia-milano-2019.xlsm data/

import json, re, sys
from pathlib import Path
import openpyxl

def clean(s):
    if s is None:
        return None
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s or None

def domande(ws):
    """Blocchi: riga con numero in A e testo in B, poi 3 righe di risposte.
    La risposta esatta è marcata 'ok' in colonna Y (indice 24).
    La X in colonna A è la risposta data dall'ultimo utente: si ignora."""
    out, cur, warn = [], None, []
    for r in ws.iter_rows(min_row=5, values_only=True):
        a, b, ok = r[0], clean(r[1]), clean(r[24])
        if a is None and b is None:
            if cur: out.append(cur); cur = None
            continue
        if a is not None and str(a).strip().isdigit() and b:
            if cur: out.append(cur)
            cur = {"n": int(a), "q": b, "a": [], "ok": None}
        elif cur is not None and b:
            if ok == "ok":
                if cur["ok"] is not None:
                    warn.append(f"domanda {cur['n']}: due risposte 'ok'")
                cur["ok"] = len(cur["a"])
            cur["a"].append(b)
    if cur: out.append(cur)
    # Correzioni a mano su errori dell'xlsm (motivo a fianco).
    OVERRIDE = {786: 0}  # due 'ok' nel foglio; la SS 33 del Sempione parte dall'Arco della Pace
    # Sezione lingua straniera: 869-949 inglese, 950-977 francese (verificato a mano sul foglio).
    for d in out:
        if d["n"] in OVERRIDE: d["ok"] = OVERRIDE[d["n"]]
        if 869 <= d["n"] <= 949: d["lang"] = "en"
        elif d["n"] >= 950: d["lang"] = "fr"
        if len(d["a"]) != 3: warn.append(f"domanda {d['n']}: {len(d['a'])} risposte")
        if d["ok"] is None: warn.append(f"domanda {d['n']}: nessuna risposta esatta")
    return out, warn

def sequenze(ws):
    """Percorsi e Confluenze: colonna A = id ('12', '12a', '12b'…), colonna B = testo.
    La riga con id numerico è il titolo, le righe con suffisso lettera sono i passi."""
    out, cur = [], None
    for r in ws.iter_rows(values_only=True):
        a, b = clean(r[0]), clean(r[1])
        if not a or not b: continue
        m = re.fullmatch(r"(\d+)([a-z]*)", a)
        if not m: continue
        n, suf = int(m.group(1)), m.group(2)
        if not suf:
            cur = {"n": n, "titolo": b, "passi": []}
            out.append(cur)
        elif cur and cur["n"] == n:
            cur["passi"].append(b)
    return out

def cosadove(ws):
    """Tre gruppi di colonne (Cosa, Dove) con intestazione di categoria in riga 1."""
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for col in (0, 3, 6):
        cat = clean(rows[0][col])
        for r in rows[2:]:
            cosa, dove = clean(r[col]), clean(r[col + 1])
            if cosa and dove:
                out.append({"cat": cat, "cosa": cosa, "dove": dove})
    return out

def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    dst = Path(dst); dst.mkdir(exist_ok=True)
    d, warn = domande(wb["Domande"])
    p = sequenze(wb["Percorsi"])
    c = sequenze(wb["Confluenze"])
    cd = cosadove(wb["CosaDove"])
    for name, obj in (("domande", d), ("percorsi", p), ("confluenze", c), ("cosadove", cd)):
        (dst / f"{name}.json").write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), "utf-8")
        print(f"{name}: {len(obj)}")
    for w in warn: print("⚠", w)
    vuoti = [x["n"] for x in p + c if not x["passi"]]
    if vuoti: print("⚠ sequenze senza passi:", vuoti)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
